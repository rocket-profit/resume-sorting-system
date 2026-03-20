import streamlit as st
import os
import json # Added for handling structured data
from dotenv import load_dotenv
from PyPDF2 import PdfReader
from langchain_google_genai import ChatGoogleGenerativeAI, GoogleGenerativeAIEmbeddings
from sklearn.metrics.pairwise import cosine_similarity
import pandas as pd
import docx
from fpdf import FPDF

# 1. SETUP & PAGE CONFIG (Must be first)
if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

if st.session_state.authenticated:
    st.set_page_config(page_title='Rocket Profit AI', page_icon="🚀", layout="wide", initial_sidebar_state="collapsed")
else:
    st.set_page_config(page_title='Rocket Profit AI', page_icon="🚀", layout="wide", initial_sidebar_state="expanded")

# 2. LOAD SECRETS (The Bridge for Cloud Deployment)
load_dotenv() 

try:
    if "GEMINI_API_KEY" in st.secrets:
        os.environ["GEMINI_API_KEY"] = st.secrets["GEMINI_API_KEY"]
    if "APP_PASSWORD" in st.secrets:
        os.environ["APP_PASSWORD"] = st.secrets["APP_PASSWORD"]
except FileNotFoundError:
    pass 

Secret_pass = os.getenv("APP_PASSWORD")

# 3. THE GATEKEEPER (Absolute Security Lock)
if not st.session_state.authenticated:
    with st.sidebar:
        st.header("Security Gateway")
        user_pass = st.text_input("Enter Access Key", type="password")
        
        if user_pass == Secret_pass and Secret_pass is not None:
            st.session_state.authenticated = True
            st.rerun()  # Instantly reloads app to collapse sidebar
        elif user_pass != "":
            st.error("Access Denied")
            
    st.info("🔒 Please enter your access key in the sidebar to unlock the AI Sorter.")
    st.stop() # CRITICAL: This completely stops the app from rendering the UI below!

st.sidebar.success("Access Granted")

# 4. INITIALIZE MODELS
@st.cache_resource
def init_models():
    api_key = os.environ.get("GEMINI_API_KEY")
    embeddings = GoogleGenerativeAIEmbeddings(
        model='gemini-embedding-2-preview', 
        google_api_key=api_key
    )
    llm = ChatGoogleGenerativeAI(
        model='gemini-2.5-flash-lite', 
        google_api_key=api_key
    )
    return embeddings, llm

embeddings, llm = init_models()

# 5. HELPER FUNCTION
def extract_text(feed):
    text = ""
    file_extension = feed.name.split('.')[-1].lower()
    try:
        if file_extension == 'pdf':
            pdf_reader = PdfReader(feed)
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted
        elif file_extension == 'docx':
            doc = docx.Document(feed)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif file_extension == 'txt':
            text = feed.getvalue().decode("utf-8")
    except Exception as e:
        st.error(f"Error reading {feed.name}: {e}")
    return text



def create_pdf_report(results):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", style="B", size=16)
    pdf.cell(200, 10, txt="AI Candidate Ranking Report", ln=True, align='C')
    pdf.ln(10)

    for res in results:
        pdf.set_font("Arial", style="B", size=12)
        pdf.cell(200, 8, txt=f"Candidate: {res['Name']} (Fit: {res['Similarity']}%)", ln=True)
        
        pdf.set_font("Arial", size=10)
        pdf.cell(200, 6, txt=f"Email: {res['Email']} | Phone: {res['Phone']}", ln=True)
        pdf.ln(4)
        
        pdf.set_font("Arial", style="B", size=10)
        pdf.cell(200, 6, txt="Pros:", ln=True)
        pdf.set_font("Arial", size=10)
        for pro in res['Analysis'].get('pros', []):
            # Encode/decode handles weird bullet point characters that crash PDFs
            safe_pro = pro.encode('latin-1', 'replace').decode('latin-1') 
            pdf.multi_cell(0, 6, txt=f"- {safe_pro}")
        
        pdf.ln(2)
        pdf.set_font("Arial", style="B", size=10)
        pdf.cell(30, 6, txt="Critical Con: ", ln=False)
        pdf.set_font("Arial", size=10)
        safe_con = res['Analysis'].get('critical_con', 'N/A').encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 6, txt=safe_con)
        
        pdf.set_font("Arial", style="B", size=10)
        pdf.cell(20, 6, txt="Verdict: ", ln=False)
        pdf.set_font("Arial", size=10)
        safe_verdict = res['Analysis'].get('verdict', 'N/A').encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 6, txt=safe_verdict)
        
        pdf.ln(8)
        
    return pdf.output(dest='S').encode('latin-1')



# 6. UI LAYOUT
st.title("Get top candidates in 2 minutes without manual screening")
st.markdown("### Powered by Rocket Profit Systems")

col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("1. Job Description")
    jd_text = st.text_area("Paste the Job Description here...", height=250)
    
with col2:
    st.subheader("2. Upload Resumes")
    uploaded_files = st.file_uploader("Upload Resumes", type=["pdf", "docx", "txt"], accept_multiple_files=True)
    
# 7. PROCESSING SECTION
if st.button("Start Sorting"):
    if not jd_text or not uploaded_files:
        st.warning("Please Provide both a Job Description and at least 1 Resume")
    else:
        # Create a dynamic placeholder
        status_text = st.empty()
        
        try:
            status_text.info("⏳ Hang on, extracting text from resumes...")
            all_extracted = [{"name": file.name, "text": extract_text(file)} for file in uploaded_files]
            resume_data = [r for r in all_extracted if r.get("text") and str(r["text"]).strip()]
            
            if not resume_data:
                status_text.error("No readable text found. Process stopped.")
                st.stop()

            status_text.info("🧠 Organizing and calculating semantic match...")
            jd_vector = embeddings.embed_query(jd_text)
            resume_texts = [r["text"] for r in resume_data]
            resume_vectors = embeddings.embed_documents(resume_texts)
            scores = cosine_similarity([jd_vector], resume_vectors)[0]

            status_text.info("📊 Ranking top candidates...")
            scored_resumes = [{"name": resume_data[i]["name"], "text": resume_texts[i], "score": score} for i, score in enumerate(scores)]
            scored_resumes = sorted(scored_resumes, key=lambda x: x["score"], reverse=True)

            MIN_SCORE, TOP_PERCENT = 0.50, 0.50
            max_candidates = max(1, int(len(scored_resumes) * TOP_PERCENT))
            final_candidates = [c for c in scored_resumes[:max_candidates] if c["score"] >= MIN_SCORE]
            if not final_candidates:
                final_candidates = [scored_resumes[0]]

            status_text.info("🤖 AI is deeply analyzing profiles... almost done!")
            results = []
            for candidate in final_candidates:
                prompt = f"""You are an expert HR. Critically evaluate this resume against the JD. 
                Resume: {candidate['text']} 
                JD: {jd_text}
                
                Provide your response STRICTLY in JSON format:
                {{
                    "candidate_name": "Full Name from resume",
                    "contact_info": {{"email": "Email address", "phone": "Phone number"}},
                    "analysis": {{"pros": ["bullet 1", "bullet 2"], "critical_con": "Main missing skill", "verdict": "1-2 line summary"}}
                }}
                """
                analysis = llm.invoke(prompt)
                try:
                    clean_content = analysis.content.replace("```json", "").replace("```", "").strip()
                    data = json.loads(clean_content)
                    results.append({
                        "Filename": candidate["name"],
                        "Name" : data.get("candidate_name", candidate["name"]),
                        "Email": data.get("contact_info", {}).get("email", "N/A"),
                        "Phone": data.get("contact_info", {}).get("phone", "N/A"),
                        "Similarity" : round(candidate["score"] * 100, 1),
                        "Analysis" : data.get("analysis", {})
                    })
                except Exception:
                    continue
            
            # Clear the loading text and show success
            status_text.empty()
            st.success(f"✅ Successfully filtered down to the top {len(results)} candidate(s)!")

            # --- DISPLAY & EXPORT ---
            for res in results:
                with st.expander(f"{res['Name']} - Fit Score: {res['Similarity']}%"):
                    # ... (Keep your existing display logic here) ...
                    st.write(f"Email: {res['Email']} | Phone: {res['Phone']}")
                    st.write(res['Analysis']) # Simplified for brevity

            if results:
                col_csv, col_pdf = st.columns(2)
                
                with col_csv:
                    # CSV Export
                    df = pd.DataFrame(results)
                    df['Pros'] = df['Analysis'].apply(lambda x: " | ".join(x.get('pros', [])))
                    df['Critical Con'] = df['Analysis'].apply(lambda x: x.get('critical_con', ''))
                    df['Verdict'] = df['Analysis'].apply(lambda x: x.get('verdict', ''))
                    export_df = df[['Name', 'Email', 'Phone', 'Similarity', 'Pros', 'Critical Con', 'Verdict']]
                    csv = export_df.to_csv(index=False).encode('utf-8')
                    st.download_button("📥 Download CSV", data=csv, file_name='candidates.csv', mime='text/csv')
                
                with col_pdf:
                    # PDF Export
                    pdf_bytes = create_pdf_report(results)
                    st.download_button("📄 Download PDF Report", data=pdf_bytes, file_name='rocket_profit_report.pdf', mime='application/pdf')

        except Exception as e:
            status_text.error(f"An error occurred: {e}")